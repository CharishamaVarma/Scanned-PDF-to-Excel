"""
Scanned PDF-to-Excel Application
Extracts text from scanned PDFs using PyMuPDF + Tesseract OCR,
then exports the recognized lines to an editable Excel workbook.
Optional: compare extracted text with an existing Excel file.
"""
from pathlib import Path
import re
import sys
import tkinter as tk
from tkinter import filedialog, messagebox

import fitz  # PyMuPDF
import pytesseract
from PIL import Image
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def configure_tesseract():
    """Use Tesseract from PATH, or common Windows installation paths."""
    candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    if not any(Path(p).exists() for p in candidates) and not shutil_which("tesseract"):
        raise RuntimeError(
            "Tesseract OCR was not found. Install Tesseract OCR and reopen this app. "
            "See README.md for instructions."
        )
    for candidate in candidates:
        if Path(candidate).exists():
            pytesseract.pytesseract.tesseract_cmd = candidate
            break


def shutil_which(command):
    import shutil
    return shutil.which(command)


def extract_pdf_lines(pdf_path, dpi=200):
    """Render each PDF page and OCR it; return one record per recognized line."""
    records = []
    pdf = fitz.open(pdf_path)
    try:
        for page_index, page in enumerate(pdf, start=1):
            pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
            image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            text = pytesseract.image_to_string(image, config="--psm 6")
            for line_index, line in enumerate(text.splitlines(), start=1):
                cleaned = re.sub(r"\s+", " ", line).strip()
                if cleaned:
                    records.append({
                        "Page": page_index,
                        "Line": line_index,
                        "Extracted Text": cleaned,
                    })
    finally:
        pdf.close()
    return records


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="17365D")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 70)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def save_results(records, output_path, source_pdf, comparison_path=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Extracted Data"
    ws.append(["Page", "Line", "Extracted Text"])
    for record in records:
        ws.append([record["Page"], record["Line"], record["Extracted Text"]])
    style_sheet(ws)

    info = wb.create_sheet("Run Info")
    info.append(["Field", "Value"])
    info.append(["Source PDF", str(source_pdf)])
    info.append(["OCR lines extracted", len(records)])
    info.append(["Comparison Excel", str(comparison_path) if comparison_path else "Not used"])
    info.append(["Note", "OCR output may contain recognition errors; review before using."])
    style_sheet(info)

    if comparison_path:
        compare_ws = wb.create_sheet("Comparison")
        compare_ws.append(["Extracted Text", "Match in Excel?", "Matching Excel Row"])
        existing = load_workbook(comparison_path, read_only=True, data_only=True)
        lookup = {}
        for sheet in existing.worksheets:
            for row_num, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                for value in row:
                    if value is not None:
                        key = re.sub(r"\s+", " ", str(value)).strip().casefold()
                        if key:
                            lookup.setdefault(key, []).append(f"{sheet.title}!Row {row_num}")
        existing.close()
        for record in records:
            value = record["Extracted Text"]
            key = re.sub(r"\s+", " ", value).strip().casefold()
            hits = lookup.get(key, [])
            compare_ws.append([value, "Yes" if hits else "No", ", ".join(hits)])
        style_sheet(compare_ws)

    wb.save(output_path)


def main():
    root = tk.Tk()
    root.withdraw()
    root.update()

    try:
        configure_tesseract()
        pdf_path = filedialog.askopenfilename(
            title="Select scanned PDF",
            filetypes=[("PDF files", "*.pdf")],
        )
        if not pdf_path:
            return

        compare = messagebox.askyesno(
            "Optional Excel comparison",
            "Do you want to compare OCR text against an existing Excel workbook?"
        )
        comparison_path = None
        if compare:
            comparison_path = filedialog.askopenfilename(
                title="Select Excel file to compare",
                filetypes=[("Excel files", "*.xlsx *.xlsm")],
            )
            if not comparison_path:
                messagebox.showinfo("Comparison skipped", "No Excel file selected. Exporting OCR results only.")

        output_path = filedialog.asksaveasfilename(
            title="Save extracted Excel workbook",
            defaultextension=".xlsx",
            initialfile=f"{Path(pdf_path).stem}_extracted.xlsx",
            filetypes=[("Excel workbook", "*.xlsx")],
        )
        if not output_path:
            return

        root.deiconify()
        root.title("Scanned PDF-to-Excel")
        root.geometry("420x120")
        label = tk.Label(root, text="Processing PDF with OCR...\nPlease wait.", padx=20, pady=25)
        label.pack()
        root.update()

        records = extract_pdf_lines(pdf_path)
        if not records:
            messagebox.showwarning(
                "No text detected",
                "OCR did not detect readable text. Try a clearer scan or higher-resolution PDF."
            )
        save_results(records, output_path, pdf_path, comparison_path)
        root.destroy()
        messagebox.showinfo(
            "Completed",
            f"Finished!\nExtracted lines: {len(records)}\nSaved Excel file:\n{output_path}"
        )
    except Exception as exc:
        try:
            root.destroy()
        except Exception:
            pass
        messagebox.showerror("Could not process PDF", str(exc) + "\n\nCheck README.md for setup instructions.")


if __name__ == "__main__":
    main()
